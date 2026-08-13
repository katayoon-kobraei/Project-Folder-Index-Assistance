"""
Tests for src/project_info/reader.py — the "LISTADO PROYECTOS POR AÑOS"
workbook reader backing the Proyectos Info page. Covers both
split_workbook_by_year() (master multi-sheet workbook -> one file per
year) and load_project_info() (reads those per-year files back).
Run with: pytest
"""

import datetime
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import openpyxl
import pytest

from src.project_info.reader import load_project_info, split_workbook_by_year

HEADERS = [
    "NUMERO DE PROYECTO", "PLANNING", "NOMBRE", "TRABAJOS A REALIZAR",
    "TITULO ENTERO DEL PROYECTO", "CIUDAD", "PROVINCIA", "FECHA DEL PROYECTO",
    "PROMOTOR", "TIPO", "SUBTIPO1", "SUBTIPO2",
    "PRESUPUESTO EN Nº\nEJECUCIÓN MATERIAL", "NOTAS", "FIRMADO",
    "FECHA DE LA FIRMA", "VISADO", "NUMERO DE EXPEDIENTE", "FECHA DEL VISADO",
    "COMIENZO DE LA OBRA", "FIN DE LA OBRA", "IMPORTE DEL CONTRATO",
    "LISTADO DE EMPLEADOS",
]


def _write_sheet(ws, rows: list[dict]) -> None:
    ws.append(HEADERS)
    for row in rows:
        ws.append([row.get(h) for h in HEADERS])


def _make_year_file(dir_path, year: str, rows: list[dict]) -> None:
    """Write dir_path/{year}.xlsx directly — bypasses split_workbook_by_year
    so load_project_info() can be tested against per-year files on its
    own, independent of the split step."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = year
    _write_sheet(ws, rows)
    Path(dir_path).mkdir(parents=True, exist_ok=True)
    wb.save(str(Path(dir_path) / f"{year}.xlsx"))


def _make_master_workbook(tmp_path, sheets: dict) -> str:
    """sheets: {sheet_name: [row_dict, ...]} — a single multi-sheet
    workbook, the shape split_workbook_by_year() takes as input."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(sheet_name)
        _write_sheet(ws, rows)
    path = str(tmp_path / "master.xlsx")
    wb.save(path)
    return path


# ---------------------------------------------------------------------
# load_project_info() — reading per-year files from a directory
# ---------------------------------------------------------------------


def test_load_project_info_maps_known_columns(tmp_path):
    _make_year_file(tmp_path, "2024", [{
        "NUMERO DE PROYECTO": "24-001-01",
        "PLANNING": "Acabado",
        "NOMBRE": "AYTO TORRENT - URB VENTETA",
        "TRABAJOS A REALIZAR": "26-001-01 Informe tecnico_f.pdf",
        "TITULO ENTERO DEL PROYECTO": "Informe técnico complementario",
        "CIUDAD": "TORRENT",
        "PROVINCIA": "VALENCIA",
        "FECHA DEL PROYECTO": datetime.datetime(2024, 6, 1),
        "PROMOTOR": "Ajuntament de Torrent.",
        "TIPO": "PROYECTO",
        "SUBTIPO1": "Establecimiento",
        "SUBTIPO2": "Consum",
        "PRESUPUESTO EN Nº\nEJECUCIÓN MATERIAL": 12345.67,
        "NOTAS": "Nota de prueba",
        "FIRMADO": "SI",
        "FECHA DE LA FIRMA": datetime.datetime(2024, 6, 5),
        "VISADO": "SI",
        "NUMERO DE EXPEDIENTE": "EXP-001",
        "FECHA DEL VISADO": datetime.datetime(2024, 6, 10),
        "COMIENZO DE LA OBRA": datetime.datetime(2024, 7, 1),
        "IMPORTE DEL CONTRATO": 99999,
        "LISTADO DE EMPLEADOS": "Juan Pérez, María López",
    }])

    rows = load_project_info(str(tmp_path))
    assert len(rows) == 1
    r = rows[0]
    assert r["year"] == 2024
    assert r["company"] == "AYTO TORRENT"
    assert r["project_name"] == "AYTO TORRENT - URB VENTETA"
    assert r["project_number"] == "24-001-01"
    assert r["status"] == "Acabado"
    assert r["work_description"] == "26-001-01 Informe tecnico_f.pdf"
    assert r["full_title"] == "Informe técnico complementario"
    assert r["location"] == "TORRENT (VALENCIA)"
    assert r["client"] == "Ajuntament de Torrent."
    assert r["work_type"] == "PROYECTO"
    assert r["project_type"] == "Establecimiento / Consum"
    assert r["budget_execution"] == "12345.67"
    assert r["notes"] == "Nota de prueba"
    assert r["signed"] == "SI"
    assert r["signed_date"] == "05/06/2024"
    assert r["visa"] == "SI"
    assert r["file_number"] == "EXP-001"
    assert r["visa_date"] == "10/06/2024"
    assert r["work_start_date"] == "01/07/2024"
    assert r["contract_amount"] == "99999"
    assert r["start_date"] == "01/06/2024"
    assert r["deadline"] is None
    assert r["employee_list"] == "Juan Pérez, María López"


def test_load_project_info_employee_list_none_when_column_absent(tmp_path):
    # Real-world case: a year file that predates the "LISTADO DE
    # EMPLEADOS" column being added — must come through as None, not
    # crash or KeyError.
    _make_year_file(tmp_path, "2025", [{"NOMBRE": "SIN COLUMNA DE EMPLEADOS"}])
    rows = load_project_info(str(tmp_path))
    assert rows[0]["employee_list"] is None


def test_load_project_info_derives_company_without_dash(tmp_path):
    _make_year_file(tmp_path, "2025", [{"NOMBRE": "SOLO NOMBRE SIN GUION"}])
    rows = load_project_info(str(tmp_path))
    assert rows[0]["company"] == "SOLO NOMBRE SIN GUION"


def test_load_project_info_groups_company_regardless_of_dash_spacing(tmp_path):
    # Real NOMBRE data is inconsistent about spacing around the "-" that
    # separates company from site — all four of these must resolve to
    # the same "PLENOIL" grouping key, not just the one with spaces on
    # both sides.
    _make_year_file(tmp_path, "2026", [
        {"NOMBRE": "PLENOIL - TETUAN 2 GANDIA"},
        {"NOMBRE": "PLENOIL- AVDA GENERALITAT MASSANASSA"},
        {"NOMBRE": "PLENOIL -CTRA DEL MIG 36 HOSPITALET"},
        {"NOMBRE": "PLENOIL-CTRA NACIONAL 340"},
    ])
    rows = load_project_info(str(tmp_path))
    assert [r["company"] for r in rows] == ["PLENOIL"] * 4


def test_load_project_info_company_only_splits_on_first_dash(tmp_path):
    # "SEINZA - CV-310 NAQUERA" has a second "-" inside the site code
    # itself (CV-310) — the company must come from the FIRST dash only.
    _make_year_file(tmp_path, "2026", [{"NOMBRE": "SEINZA - CV-310 NAQUERA"}])
    rows = load_project_info(str(tmp_path))
    assert rows[0]["company"] == "SEINZA"


def test_load_project_info_skips_blank_spacer_rows(tmp_path):
    _make_year_file(tmp_path, "2026", [
        {"NOMBRE": "PLENOIL - SITE A"},
        {},  # blank spacer row — no NOMBRE
        {"NOMBRE": "PLENOIL - SITE B"},
    ])
    rows = load_project_info(str(tmp_path))
    assert [r["project_name"] for r in rows] == ["PLENOIL - SITE A", "PLENOIL - SITE B"]


def test_load_project_info_preserves_sheet_row_order_not_alphabetical(tmp_path):
    # Multiple rows can share the same NOMBRE (different documents/tasks
    # for the same project) — order must follow the sheet, NOT be
    # re-sorted alphabetically, since a later row can depend on an
    # earlier one being shown first.
    _make_year_file(tmp_path, "2026", [
        {"NOMBRE": "ZETA PROJECT", "PLANNING": None},
        {"NOMBRE": "AYTO TORRENT - URB VENTETA", "PLANNING": None},
        {"NOMBRE": "AYTO TORRENT - URB VENTETA", "PLANNING": "Acabado"},
    ])
    rows = load_project_info(str(tmp_path))
    assert [r["project_name"] for r in rows] == [
        "ZETA PROJECT", "AYTO TORRENT - URB VENTETA", "AYTO TORRENT - URB VENTETA",
    ]
    assert [r["status"] for r in rows] == [None, None, "Acabado"]


def test_load_project_info_ignores_non_project_files(tmp_path):
    _make_year_file(tmp_path, "2024", [{"NOMBRE": "REAL ROW"}])
    # A stray unrelated xlsx sitting in the same folder (e.g. a leftover
    # download) that isn't named 2024/2025/2026.xlsx must be ignored.
    (tmp_path / "notes.xlsx").write_bytes(b"not a real workbook")
    rows = load_project_info(str(tmp_path))
    assert [r["project_name"] for r in rows] == ["REAL ROW"]


def test_load_project_info_coerces_stray_non_text_values(tmp_path):
    # A handful of real rows have a stray date typed into a normally-text
    # column (e.g. SUBTIPO1) — must not crash, just come through as text.
    _make_year_file(tmp_path, "2024", [{
        "NOMBRE": "X - Y",
        "SUBTIPO1": datetime.datetime(2021, 1, 1),
    }])
    rows = load_project_info(str(tmp_path))
    assert rows[0]["project_type"] == "01/01/2021"


def test_load_project_info_ids_are_unique_and_ordered_by_year(tmp_path):
    _make_year_file(tmp_path, "2024", [{"NOMBRE": "A"}, {"NOMBRE": "B"}])
    _make_year_file(tmp_path, "2025", [{"NOMBRE": "C"}])
    rows = load_project_info(str(tmp_path))
    assert [r["id"] for r in rows] == [0, 1, 2]
    assert [(r["year"], r["project_name"]) for r in rows] == [
        (2024, "A"), (2024, "B"), (2025, "C"),
    ]


def test_load_project_info_skips_missing_year_files_without_erroring(tmp_path):
    # Only 2024.xlsx exists — 2025/2026 simply aren't there yet. Must
    # still return what IS available, not error out entirely.
    _make_year_file(tmp_path, "2024", [{"NOMBRE": "ONLY 2024"}])
    rows = load_project_info(str(tmp_path))
    assert [r["project_name"] for r in rows] == ["ONLY 2024"]


def test_load_project_info_missing_directory_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        load_project_info(str(tmp_path / "does_not_exist"))


def test_load_project_info_empty_directory_raises(tmp_path):
    # Directory exists but has none of the 3 expected year files.
    with pytest.raises(FileNotFoundError):
        load_project_info(str(tmp_path))


# ---------------------------------------------------------------------
# split_workbook_by_year() — master workbook -> one file per year
# ---------------------------------------------------------------------


def test_split_workbook_by_year_writes_one_file_per_year_with_data(tmp_path):
    master = _make_master_workbook(tmp_path, {
        "2023": [{"NOMBRE": "SHOULD NOT BE SPLIT"}],  # not in YEARS_WITH_DATA
        "2024": [{"NOMBRE": "A"}, {"NOMBRE": "B"}],
        "2025": [{"NOMBRE": "C"}],
        # 2026 sheet absent entirely — must simply be skipped, not error
    })
    output_dir = tmp_path / "split"
    written = split_workbook_by_year(master, str(output_dir))

    written_names = {Path(p).name for p in written}
    assert written_names == {"2024.xlsx", "2025.xlsx"}
    assert not (output_dir / "2023.xlsx").exists()
    assert not (output_dir / "2026.xlsx").exists()


def test_split_workbook_by_year_output_reads_back_correctly(tmp_path):
    master = _make_master_workbook(tmp_path, {
        "2026": [
            {"NOMBRE": "FIRST ROW", "PLANNING": None},
            {"NOMBRE": "SECOND ROW", "PLANNING": "Acabado",
             "FECHA DEL PROYECTO": datetime.datetime(2026, 6, 1)},
        ],
    })
    output_dir = tmp_path / "split"
    split_workbook_by_year(master, str(output_dir))

    rows = load_project_info(str(output_dir))
    assert [r["project_name"] for r in rows] == ["FIRST ROW", "SECOND ROW"]
    assert rows[1]["status"] == "Acabado"
    assert rows[1]["start_date"] == "01/06/2026"


def test_split_workbook_by_year_is_rerunnable(tmp_path):
    # Running it twice (e.g. after the master workbook was updated) must
    # just overwrite the same 3 files, not fail or duplicate anything.
    master = _make_master_workbook(tmp_path, {"2024": [{"NOMBRE": "V1"}]})
    output_dir = tmp_path / "split"
    split_workbook_by_year(master, str(output_dir))

    master2 = _make_master_workbook(tmp_path, {"2024": [{"NOMBRE": "V2"}]})
    written = split_workbook_by_year(master2, str(output_dir))
    assert len(written) == 1

    rows = load_project_info(str(output_dir))
    assert [r["project_name"] for r in rows] == ["V2"]


def test_split_workbook_by_year_missing_source_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        split_workbook_by_year(str(tmp_path / "nope.xlsx"), str(tmp_path / "out"))
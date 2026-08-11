"""
Writes edits made in the Proyectos Info detail page's "Editar" form, and
brand-new rows added via "Nuevo proyecto", back into the per-year
project_info_dir files (2024.xlsx / 2025.xlsx / 2026.xlsx — see
reader.py's module docstring).

Deliberately narrow scope, unlike split_workbook_by_year() in reader.py:
neither function here ever rewrites the whole sheet. update_project_info
_row() opens the workbook normally (not read_only), touches only the
specific cells being edited on one already-existing row, and saves —
every other row, every other column, and any formatting/data validation
openpyxl doesn't need to touch stays exactly as it was. `row_number`
(see reader.py's _row_dict_from_raw) is what makes this possible: it's
the row's real 1-based position in the sheet, captured at read time, so
there's no need to re-search for the row by content (which would break
as soon as two rows share a NOMBRE). append_project_info_row() is the
same idea in reverse: it only ever writes to ONE brand-new row, one past
whatever the sheet's current last row is — every existing row is
untouched.

This is the one place in the whole app that writes to a file inside
project_info_dir — everywhere else (the crawled P: index itself) stays
strictly read-only. Two things worth knowing about that:
  1. project_info_dir commonly points inside P:\\ (see config.yaml) — so
     yes, this does modify a file that lives on the P: drive. That's a
     deliberate exception for this one hand-maintained dataset, not an
     oversight of the read-only rule for the crawled archive.
  2. If project_info_dir's files are ever regenerated from the master
     workbook via split_workbook_by_year() / scripts/split_project_info
     _by_year.py, that overwrites the whole file from the master again —
     any edit (or new row) made through the app that was never carried
     back into the master workbook would be lost. Worth keeping the
     master workbook itself updated too, or treating the per-year files
     as the source of truth once editing through the app starts.
"""

from __future__ import annotations

import datetime
from pathlib import Path

import openpyxl

from .reader import KEY_TO_HEADER, _KEY_IS_DATE


def _parse_date(value: str) -> datetime.date | None:
    """Inverse of reader._format_date: 'dd/mm/yyyy' -> a real date, so
    the cell keeps behaving like a date in Excel (sorting, filters,
    etc.) instead of turning into a text string that merely looks like
    one. Raises ValueError with a Spanish message on anything that isn't
    a valid dd/mm/yyyy string — the caller is expected to show that to
    the user rather than silently writing bad data."""
    value = (value or "").strip()
    if not value:
        return None
    try:
        return datetime.datetime.strptime(value, "%d/%m/%Y").date()
    except ValueError:
        raise ValueError(
            f"Fecha no válida: «{value}». Usa el formato día/mes/año, por ejemplo 05/06/2024."
        )


def _prepare_values(values: dict) -> dict[str, object]:
    """updates/values (field key -> raw string from a form) -> the same
    keys mapped to their actual cell-ready values (a real date for date
    fields, a stripped string-or-None for everything else). Shared by
    update_project_info_row() and append_project_info_row() so the two
    can never validate/coerce differently. Keys not in
    reader.KEY_TO_HEADER are dropped silently — callers can pass a whole
    form's worth of fields without filtering out read-only/derived ones
    (company, location, project_type, id, row_number, year) themselves."""
    parsed: dict[str, object] = {}
    for key, value in values.items():
        header = KEY_TO_HEADER.get(key)
        if header is None:
            continue
        if _KEY_IS_DATE.get(key):
            parsed[key] = _parse_date(value)
        else:
            value = (value or "").strip() if isinstance(value, str) else value
            parsed[key] = value or None
    return parsed


def _header_to_col(ws) -> dict:
    header_row = next(ws.iter_rows(min_row=1, max_row=1))
    return {
        (cell.value.strip() if isinstance(cell.value, str) else cell.value): cell.column
        for cell in header_row
        if cell.value
    }


def _open_year_file(dir_path: str, year: int):
    file_path = Path(dir_path) / f"{year}.xlsx"
    if not file_path.exists():
        raise FileNotFoundError(f"No se encuentra el archivo de {year}: {file_path}")
    return file_path, openpyxl.load_workbook(str(file_path))


def update_project_info_row(dir_path: str, year: int, row_number: int, updates: dict) -> None:
    """Write `updates` (field key -> new value, plain strings — empty
    string clears the cell) into row `row_number` of `{dir_path}/
    {year}.xlsx`.

    Raises FileNotFoundError if the year file doesn't exist, and
    ValueError if a date field's value isn't a valid dd/mm/yyyy string."""
    # Validated BEFORE opening the workbook for writing, so a typo in
    # one field can't leave the file half-edited.
    parsed = _prepare_values(updates)

    file_path, wb = _open_year_file(dir_path, year)
    try:
        ws = wb.worksheets[0]
        header_to_col = _header_to_col(ws)

        for key, value in parsed.items():
            header = KEY_TO_HEADER[key]
            col = header_to_col.get(header)
            if col is None:
                continue  # this file's sheet doesn't have that column at all
            # NOT ws.cell(row, col, value=value) — openpyxl's cell()
            # treats value=None as "no value passed" and leaves the
            # existing content alone, so clearing a field silently did
            # nothing. Getting the cell first and assigning .value
            # directly clears it correctly.
            ws.cell(row=row_number, column=col).value = value

        wb.save(str(file_path))
    finally:
        wb.close()


def append_project_info_row(dir_path: str, year: int, values: dict) -> int:
    """Add a brand-new row at the end of `{dir_path}/{year}.xlsx`,
    writing `values` (same field-key -> value shape as
    update_project_info_row) into the matching columns — any column not
    present in `values`, or not part of this particular sheet at all, is
    simply left blank. Returns the new row's 1-based row_number.

    Raises FileNotFoundError if the year file doesn't exist, ValueError
    if a date field isn't a valid dd/mm/yyyy string, and ValueError if
    `values["project_name"]` (NOMBRE) is blank — every row needs a name
    to ever be found again by load_project_info()."""
    if not (values.get("project_name") or "").strip():
        raise ValueError("El proyecto necesita un nombre (NOMBRE) para poder guardarse.")

    parsed = _prepare_values(values)

    file_path, wb = _open_year_file(dir_path, year)
    try:
        ws = wb.worksheets[0]
        header_to_col = _header_to_col(ws)
        row_number = ws.max_row + 1

        for key, value in parsed.items():
            header = KEY_TO_HEADER[key]
            col = header_to_col.get(header)
            if col is None:
                continue
            ws.cell(row=row_number, column=col).value = value

        wb.save(str(file_path))
    finally:
        wb.close()
    return row_number